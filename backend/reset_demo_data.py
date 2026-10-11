"""
Resets the demo environment to a clean baseline before running the eval
suite, so results are reproducible and not polluted by earlier manual
testing (uploaded custom agreements, sellers added via the UI, etc).

This restores seller_master.xlsx to the original 5 sellers, and wipes the
chroma_store vector database (which holds both the general policy AND any
custom seller agreements uploaded during testing). The general policy gets
re-indexed automatically the next time the backend starts.

Run this before eval_suite.py if you want a clean, comparable baseline.
"""

import os
import shutil
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
SELLER_MASTER_FILE = os.path.join(DATA_DIR, "seller_master.xlsx")
CHROMA_DIR = os.path.join(BASE_DIR, "chroma_store")

ORIGINAL_SELLERS = [
    ("SLR001", "Myntra"),
    ("SLR002", "ABC Fashion"),
    ("SLR003", "XYZ Traders"),
    ("SLR004", "Zenith Retail"),
    ("SLR005", "Kismat Store"),
]


def reset_seller_master():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sellers"
    ws.append(["Seller ID", "Seller Name"])

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    data_font = Font(name="Arial")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for seller_id, seller_name in ORIGINAL_SELLERS:
        ws.append([seller_id, seller_name])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18

    wb.save(SELLER_MASTER_FILE)
    print(f"Reset {SELLER_MASTER_FILE} to the original 5 sellers.")


def reset_vector_store():
    if os.path.exists(CHROMA_DIR):
        shutil.rmtree(CHROMA_DIR)
        print(f"Removed {CHROMA_DIR} (general policy will be re-indexed on next backend start).")
    else:
        print("No chroma_store found — nothing to remove.")


if __name__ == "__main__":
    reset_seller_master()
    reset_vector_store()
    print("\nDemo environment reset. Restart the backend (or run rag_engine.py) to re-index the general policy.")